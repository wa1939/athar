"""Extract TMOAP v5 workbook → Athar seed CSVs.

Reads `Budget Tracking Tool - TMOAP v5.xlsx` and writes:
  app/src/seeded/assets/seed/categories.csv
  app/src/seeded/assets/seed/transactions.csv
  app/src/seeded/assets/seed/budget_targets.csv
  app/src/seeded/assets/seed/wishlist.csv
  app/src/seeded/assets/seed/investments.csv

These get bundled only when the build is invoked with `-Pathar.seed=true`
(see app/build.gradle.kts). The DataSeeder loads them on first launch.

The CSVs are gitignored — they contain personal financial data.
"""
import csv
import datetime
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Budget Tracking Tool - TMOAP v5.xlsx"
OUT = ROOT / "app" / "src" / "seeded" / "assets" / "seed"
OUT.mkdir(parents=True, exist_ok=True)


def fmt_date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return ""


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # ── Categories ────────────────────────────────────────────────────────────
    ws = wb["Category Setup"]
    cats = []
    for r in range(3, 90):
        v = ws.cell(row=r, column=2).value
        if v and isinstance(v, str):
            cats.append(v.strip())
    with (OUT / "categories.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["name_en", "name_ar"])
        for c in cats:
            w.writerow([c, ""])
    print(f"categories.csv: {len(cats)} rows")

    # ── Transactions (Expenses + Income merged into Athar's schema) ───────────
    rows = []
    for sheet, tx_type in [("Expenses", "EXPENSE"), ("Income", "INCOME")]:
        ws = wb[sheet]
        for r in range(8, ws.max_row + 1):
            d = ws.cell(row=r, column=2).value
            vendor = ws.cell(row=r, column=3).value
            amount = ws.cell(row=r, column=4).value
            category = ws.cell(row=r, column=5).value
            notes = ws.cell(row=r, column=6).value
            if not (isinstance(d, (datetime.datetime, datetime.date)) and isinstance(amount, (int, float))):
                continue
            rows.append({
                "date": fmt_date(d),
                "vendor": str(vendor or "").strip(),
                "amount": f"{float(amount):.2f}",
                "category": str(category or "").strip(),
                "type": tx_type,
                "notes": str(notes or "").strip(),
            })
    rows.sort(key=lambda r: r["date"])
    with (OUT / "transactions.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["date", "vendor", "amount", "category", "type", "notes"])
        w.writeheader()
        w.writerows(rows)
    print(f"transactions.csv: {len(rows)} rows")

    # ── Budget Targets ────────────────────────────────────────────────────────
    ws = wb["Budget Targets"]
    targets = []
    for r in range(7, 90):
        cat = ws.cell(row=r, column=2).value
        target = ws.cell(row=r, column=7).value
        if cat and isinstance(cat, str) and isinstance(target, (int, float)) and target > 0:
            targets.append({"category": cat.strip(), "monthly_target": f"{float(target):.2f}"})
    with (OUT / "budget_targets.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["category", "monthly_target"])
        w.writeheader()
        w.writerows(targets)
    print(f"budget_targets.csv: {len(targets)} rows")

    # ── Wishlist ──────────────────────────────────────────────────────────────
    ws = wb["Wishlist"]
    wish = []
    # data starts at row 11 (1-indexed); item, cost, current_saved, ...
    for r in range(11, ws.max_row + 1):
        item = ws.cell(row=r, column=1).value
        cost = ws.cell(row=r, column=2).value
        saved = ws.cell(row=r, column=3).value
        start = ws.cell(row=r, column=11).value
        if not (item and isinstance(item, str) and isinstance(cost, (int, float))):
            continue
        wish.append({
            "name": item.strip(),
            "cost": f"{float(cost):.2f}",
            "current_saved": f"{float(saved or 0):.2f}",
            "start_month": fmt_date(start) if start else "",
        })
    with (OUT / "wishlist.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["name", "cost", "current_saved", "start_month"])
        w.writeheader()
        w.writerows(wish)
    print(f"wishlist.csv: {len(wish)} rows")

    # ── Family Investments ────────────────────────────────────────────────────
    ws = wb["استثمارات العائلة"]
    inv = []
    for r in range(2, ws.max_row + 1):
        amount = ws.cell(row=r, column=1).value
        owner = ws.cell(row=r, column=2).value
        ret = ws.cell(row=r, column=5).value
        period = ws.cell(row=r, column=6).value
        if not (isinstance(amount, (int, float)) and owner and isinstance(owner, str)):
            continue
        inv.append({
            "owner": owner.strip(),
            "amount": f"{float(amount):.2f}",
            "return": f"{float(ret):.2f}" if isinstance(ret, (int, float)) else "",
            "period": str(period or "").strip(),
        })
    with (OUT / "investments.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["owner", "amount", "return", "period"])
        w.writeheader()
        w.writerows(inv)
    print(f"investments.csv: {len(inv)} rows")

    print(f"\nAll seed CSVs written to: {OUT}")


if __name__ == "__main__":
    main()
